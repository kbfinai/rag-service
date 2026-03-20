"""Microsoft Word DOCX file parser."""

import io

from fastapi import UploadFile

from rag_service.services.parsers.base import BaseFileParser, ParseResult


class DocxParser(BaseFileParser):
    """Parser for Microsoft Word DOCX files."""

    @property
    def supported_extensions(self) -> list[str]:
        return [".docx", ".DOCX"]

    @property
    def supported_mime_types(self) -> list[str]:
        return [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ]

    async def parse(self, file: UploadFile) -> ParseResult:
        """Parse DOCX and extract text."""
        from docx import Document

        content = await self._read_file_content(file)
        docx_file = io.BytesIO(content)

        try:
            doc = Document(docx_file)

            # Extract text from paragraphs
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            # Extract text from tables
            table_texts = []
            for table in doc.tables:
                table_text = self._extract_table_text(table)
                if table_text:
                    table_texts.append(table_text)

            # Combine all text
            all_text = "\n\n".join(paragraphs)
            if table_texts:
                all_text += "\n\n" + "\n\n".join(table_texts)

            clean_text = self._clean_text(all_text)

            return ParseResult(
                content=clean_text,
                file_type="docx",
                word_count=self._count_words(clean_text),
                char_count=len(clean_text),
                metadata={
                    "filename": file.filename,
                    "file_size": len(content),
                    "paragraph_count": len(paragraphs),
                    "table_count": len(doc.tables),
                    "core_properties": self._extract_core_properties(doc),
                },
            )
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {e}") from e

    def _extract_table_text(self, table) -> str:
        """Extract text from a table."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        return "\n".join(rows)

    def _extract_core_properties(self, doc) -> dict:
        """Extract document core properties."""
        props = {}
        try:
            core = doc.core_properties
            if core.title:
                props["title"] = core.title
            if core.author:
                props["author"] = core.author
            if core.subject:
                props["subject"] = core.subject
            if core.created:
                props["created"] = str(core.created)
            if core.modified:
                props["modified"] = str(core.modified)
        except Exception:
            pass
        return props
